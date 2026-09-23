"""
Unit tests for Microsoft Word (.docx) tender document parsing.
"""
import os
import tempfile
import docx
from app.services.document_parser import get_document_parser, DocxParser, NormalizedDocument


def test_docx_parser_extracts_headings_paragraphs_and_tables():
    # Create a temporary .docx file
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp_path = tmp.name

    try:
        doc = docx.Document()
        doc.add_heading("Section 4: Technical Specifications for High Voltage Switchgear", level=1)
        doc.add_paragraph("All switchgear must conform to IS 13118 and IEC 62271-100.")
        doc.add_paragraph("Rated short-circuit breaking current shall be at least 40 kA for 3 seconds.")
        
        # Add a table with technical parameters
        table = doc.add_table(rows=1, cols=3)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Parameter"
        hdr_cells[1].text = "Required Value"
        hdr_cells[2].text = "Governing Standard"

        row = table.add_row().cells
        row[0].text = "Rated Voltage"
        row[1].text = "33 kV"
        row[2].text = "IS 13118"

        doc.save(tmp_path)

        parser = get_document_parser("application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        assert isinstance(parser, DocxParser)

        normalized = parser.parse(tmp_path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        assert isinstance(normalized, NormalizedDocument)
        assert len(normalized.blocks) >= 3

        # Verify heading extraction
        heading_block = normalized.blocks[0]
        assert "Technical Specifications" in heading_block.text
        assert heading_block.section_heading == "Section 4: Technical Specifications for High Voltage Switchgear"

        # Verify paragraph extraction
        p_block = normalized.blocks[1]
        assert "IS 13118" in p_block.text
        assert p_block.start_offset >= 0
        assert p_block.end_offset > p_block.start_offset

        # Verify table extraction
        table_block = [b for b in normalized.blocks if "33 kV" in b.text]
        assert len(table_block) == 1
        assert "Rated Voltage" in table_block[0].text

        assert "Section 4" in normalized.full_text
        assert "33 kV" in normalized.full_text

    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)
